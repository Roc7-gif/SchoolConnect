import base64
import io
import json
import os

import openpyxl
import pdfplumber
import pypdfium2 as pdfium
from django.conf import settings
from openai import OpenAI
from rest_framework.exceptions import ValidationError

from .matching import build_plan
from .parsers import _cell_str

MODEL = 'claude-sonnet-5'
MAX_CHARS = 400000  # limite d'entrée (le modèle a une fenêtre de 1M tokens, large marge)
BATCH_SIZE = 100  # lignes par appel — reste confortablement sous le plafond de sortie du modèle
MAX_TOKENS_PER_BATCH = 24000
IMAGE_SCALE = 2.5  # ~180 DPI — compromis lisibilité OCR / taille du payload
IMAGE_BATCH_SIZE = 5  # pages scannées par appel IA
MAX_IMAGE_PAGES = 30  # au-delà, le payload devient trop lourd / lent pour un seul import

# Schéma de l'outil d'extraction, au format JSON Schema (repris tel quel dans
# le tool OpenAI-compatible ci-dessous).
ROWS_SCHEMA = {
    'type': 'object',
    'properties': {
        'rows': {
            'type': 'array',
            'items': {
                'type': 'object',
                'properties': {
                    'last_name': {'type': 'string'},
                    'first_name': {'type': 'string'},
                    'classe_name': {'type': 'string', 'description': 'Nom de la classe tel que dans le fichier'},
                    'sexe': {'type': 'string', 'enum': ['M', 'F', '']},
                    'date_of_birth': {'type': 'string', 'description': 'Format YYYY-MM-DD, ou vide'},
                    'matricule': {'type': 'string'},
                    'parent_last_name': {'type': 'string'},
                    'parent_first_name': {'type': 'string'},
                    'parent_phone': {'type': 'string'},
                    'parent_relationship': {'type': 'string', 'enum': ['MERE', 'PERE', 'TUTEUR', 'AUTRE']},
                },
                'required': ['last_name', 'first_name'],
            },
        },
    },
    'required': ['rows'],
}

# Format "tool-calling" OpenAI — utilisé car l'import IA passe actuellement
# par le proxy tiers HighwayAPI (compatible OpenAI), pas par l'API Anthropic
# native. Voir parse_with_ai() : AI_PROXY_API_KEY / AI_PROXY_BASE_URL.
EXTRACT_ROWS_TOOL = {
    'type': 'function',
    'function': {
        'name': 'submit_rows',
        'description': "Soumet les lignes élèves normalisées extraites du fichier",
        'parameters': ROWS_SCHEMA,
    },
}


def _read_excel_rows(file_obj):
    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True)
    except Exception as exc:
        raise ValidationError({'file': f'Fichier illisible : {exc}'})

    sheet = wb.active
    rows_iter = sheet.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise ValidationError({'file': 'Le fichier est vide.'})

    header = [(_cell_str(h) or '') for h in header_row]
    data_rows = [
        [_cell_str(v) or '' for v in row]
        for row in rows_iter
        if row is not None and not all(v is None for v in row)
    ]
    return header, data_rows


def _render_pdf_pages_as_images(file_obj):
    """Rend chaque page du PDF en PNG (base64) pour l'envoyer à l'IA en mode vision."""
    try:
        file_obj.seek(0)
    except Exception:
        pass
    try:
        pdf = pdfium.PdfDocument(file_obj)
    except Exception as exc:
        raise ValidationError({'file': f'Fichier illisible : {exc}'})

    n_pages = len(pdf)
    if n_pages > MAX_IMAGE_PAGES:
        raise ValidationError({
            'file': (
                f'Ce PDF scanné comporte {n_pages} pages, au-delà de la limite de '
                f'{MAX_IMAGE_PAGES} pages pour le mode IA. Scindez le fichier en plusieurs '
                'imports ou utilisez le mode standard.'
            ),
        })

    images_b64 = []
    for page in pdf:
        bitmap = page.render(scale=IMAGE_SCALE)
        pil_image = bitmap.to_pil().convert('RGB')
        buf = io.BytesIO()
        pil_image.save(buf, format='PNG')
        images_b64.append(base64.b64encode(buf.getvalue()).decode('ascii'))
    return images_b64


def _extract_content(file_obj):
    """Retourne ('table', header, data_rows), ('text', contenu, None) ou ('image', pages_b64, None)."""
    filename = getattr(file_obj, 'name', '') or ''
    ext = os.path.splitext(filename)[1].lower()

    if ext != '.pdf':
        header, data_rows = _read_excel_rows(file_obj)
        return 'table', header, data_rows

    try:
        with pdfplumber.open(file_obj) as pdf:
            all_rows = []
            for page in pdf.pages:
                for table in page.extract_tables():
                    all_rows.extend(table)
            if all_rows:
                header = [_cell_str(h) or '' for h in all_rows[0]]
                data_rows = [[_cell_str(v) or '' for v in row] for row in all_rows[1:]]
                return 'table', header, data_rows

            text_parts = [page.extract_text() or '' for page in pdf.pages]
            raw_text = '\n'.join(text_parts).strip()
    except Exception as exc:
        raise ValidationError({'file': f'Fichier illisible : {exc}'})

    if not raw_text:
        # Pas de texte extractible : probablement un PDF scanné — on bascule
        # sur le mode vision (l'IA lit directement les images des pages).
        images_b64 = _render_pdf_pages_as_images(file_obj)
        return 'image', images_b64, None
    return 'text', raw_text, None


def _run_extraction(client, messages):
    response = client.chat.completions.create(
        model=MODEL,
        max_tokens=MAX_TOKENS_PER_BATCH,
        tools=[EXTRACT_ROWS_TOOL],
        tool_choice={'type': 'function', 'function': {'name': 'submit_rows'}},
        messages=messages,
    )

    tool_calls = response.choices[0].message.tool_calls or []
    call = next((c for c in tool_calls if c.function.name == 'submit_rows'), None)
    if call is None:
        raise ValidationError({'file': "L'IA n'a pas pu produire de résultat exploitable pour ce fichier."})
    try:
        args = json.loads(call.function.arguments)
    except (TypeError, ValueError):
        raise ValidationError({'file': "L'IA n'a pas pu produire de résultat exploitable pour ce fichier."})
    return args.get('rows', [])


def _call_ai(client, content_description, content):
    prompt = f"""Tu reçois un extrait d'un fichier fourni par une école (élèves/classes/parents), avec des noms de colonnes possiblement différents, dans le désordre, en français ou en anglais. Format du contenu : {content_description}.

Ta tâche : extraire chaque ligne élève via l'outil `submit_rows`, avec les champs normalisés (nom, prénom, classe, sexe, date de naissance, matricule, et les informations du parent).

Règles :
- N'invente jamais une donnée absente du fichier — laisse le champ vide plutôt que de deviner.
- Une ligne = un élève. Si plusieurs parents sont listés pour un élève, garde le contact principal.
- Recopie le nom de la classe tel qu'il apparaît dans le fichier (le rapprochement avec les classes existantes de l'école se fait automatiquement ensuite, pas besoin de t'en soucier).

Contenu extrait du fichier :
{content}
"""
    return _run_extraction(client, [{'role': 'user', 'content': prompt}])


def _call_ai_vision(client, content_description, images_b64):
    prompt = f"""Tu reçois des pages scannées (images) d'un fichier fourni par une école (élèves/classes/parents), avec des noms de colonnes possiblement différents, dans le désordre, en français ou en anglais. Format du contenu : {content_description}.

Ta tâche : lis les images (OCR) et extrais chaque ligne élève via l'outil `submit_rows`, avec les champs normalisés (nom, prénom, classe, sexe, date de naissance, matricule, et les informations du parent).

Règles :
- N'invente jamais une donnée absente du fichier — laisse le champ vide plutôt que de deviner.
- Une ligne = un élève. Si plusieurs parents sont listés pour un élève, garde le contact principal.
- Recopie le nom de la classe tel qu'il apparaît dans le fichier (le rapprochement avec les classes existantes de l'école se fait automatiquement ensuite, pas besoin de t'en soucier).
- Si un champ est illisible ou ambigu sur l'image, laisse-le vide plutôt que de deviner.
"""
    content = [{'type': 'text', 'text': prompt}]
    for img_b64 in images_b64:
        content.append({
            'type': 'image_url',
            'image_url': {'url': f'data:image/png;base64,{img_b64}'},
        })
    return _run_extraction(client, [{'role': 'user', 'content': content}])


def _normalize_row(r):
    return {
        'last_name': r.get('last_name') or '',
        'first_name': r.get('first_name') or '',
        'classe_name': r.get('classe_name') or '',
        'sexe': (r.get('sexe') or '').upper()[:1],
        'date_of_birth': r.get('date_of_birth') or '',
        'matricule': r.get('matricule') or '',
        'parent_last_name': r.get('parent_last_name') or '',
        'parent_first_name': r.get('parent_first_name') or '',
        'parent_phone': r.get('parent_phone') or '',
        'parent_relationship': r.get('parent_relationship') or 'TUTEUR',
    }


def parse_with_ai(file_obj, school, annee):
    # NOTE (temporaire) : passe par le proxy tiers HighwayAPI (compatible
    # OpenAI), pas par l'API Anthropic officielle — en attendant une vraie
    # clé ANTHROPIC_API_KEY. À remplacer dès qu'une clé officielle est
    # disponible : ce proxy n'est pas vérifié et ne doit pas recevoir de
    # données réelles d'élèves/parents en production.
    api_key = getattr(settings, 'AI_PROXY_API_KEY', None)
    if not api_key:
        raise ValidationError({
            'file': (
                "Le mode IA n'est pas encore configuré (clé API manquante côté serveur). "
                'Utilisez le mode standard en attendant.'
            ),
        })

    kind, a, b = _extract_content(file_obj)
    client = OpenAI(base_url=settings.AI_PROXY_BASE_URL, api_key=api_key)

    all_rows = []
    if kind == 'table':
        header, data_rows = a, b
        if not data_rows:
            raise ValidationError({'file': 'Le fichier est vide.'})
        table_desc = 'tableau (colonnes séparées par des tabulations, première ligne = en-têtes)'
        for i in range(0, len(data_rows), BATCH_SIZE):
            batch = data_rows[i:i + BATCH_SIZE]
            content = '\t'.join(header) + '\n' + '\n'.join('\t'.join(row) for row in batch)
            all_rows.extend(_call_ai(client, table_desc, content))
    elif kind == 'image':
        images_b64 = a
        if not images_b64:
            raise ValidationError({'file': 'Le fichier est vide.'})
        image_desc = 'pages scannées d\'un document (une image par page, envoyées par lots)'
        for i in range(0, len(images_b64), IMAGE_BATCH_SIZE):
            batch = images_b64[i:i + IMAGE_BATCH_SIZE]
            all_rows.extend(_call_ai_vision(client, image_desc, batch))
    else:
        text = a[:MAX_CHARS]
        text_desc = 'texte brut extrait du PDF, pas structuré en colonnes — à toi de repérer les enregistrements'
        all_rows.extend(_call_ai(client, text_desc, text))

    rows = [_normalize_row(r) for r in all_rows]
    return build_plan(rows, school, annee)
